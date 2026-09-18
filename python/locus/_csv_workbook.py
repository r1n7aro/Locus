"""A real openpyxl workbook backed by one CSV and its Locus view companion."""
from __future__ import annotations

import copy
import hashlib
import json
import os
import re
from pathlib import PurePosixPath

from openpyxl import Workbook
from openpyxl.cell.cell import Cell, MergedCell
from openpyxl.styles import Alignment, Color, Font
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.utils.cell import coordinate_to_tuple, get_column_letter, column_index_from_string

from ._csv import CsvViewSnapshot
from ._csv_codec import CsvSource, text_value
from ._csv_excel import COMPONENTS, apply_style, components, decode_component, encode_component, to_color


class CsvCell(Cell):
    def check_string(self, value):
        # CSV has no Excel 32,767-character limit or XML control-character rules.
        if value is None: return None
        value = str(value)
        if "\0" in value: raise ValueError("CSV cannot contain NUL bytes")
        return value


class CsvWorksheet(Worksheet):
    def __init__(self, parent, title):
        super().__init__(parent, title)
        self._csv_shadow = {}
        self._csv_extent = (1, 1)
        self._csv_appended_extent = 0

    def _get_cell(self, row, column):
        rows, columns = max(row, self._csv_extent[0]), max(column, self._csv_extent[1])
        if not 1 <= row <= 500000 or not 1 <= column <= 10000 or rows * columns > 500000:
            raise ValueError("CSV supports at most 500,000 rectangular cells and 10,000 columns")
        self._csv_extent = (rows, columns)
        if (row, column) not in self._cells:
            self._add_cell(CsvCell(self, row, column))
        return self._cells[row, column]

    def append(self, iterable):
        if isinstance(iterable, dict):
            entries = [(column_index_from_string(c) if isinstance(c, str) else c, value) for c, value in iterable.items()]
        elif not isinstance(iterable, (str, bytes)):
            entries = list(enumerate(iterable, 1))
        else: raise TypeError("append expects a row iterable or a column/value dictionary")
        row = self._current_row + 1
        for column, value in entries: self.cell(row, column, value)
        self._current_row = row
        self._csv_appended_extent = max(self._csv_appended_extent, row)

    def __delitem__(self, key):
        cell = self[key]
        if isinstance(cell, MergedCell): raise ValueError("Unmerge the range before clearing a covered cell")
        cell.value = None
        cell._style = None

    def merge_cells(self, range_string=None, start_row=None, start_column=None, end_row=None, end_column=None):
        from openpyxl.worksheet.cell_range import CellRange
        target = CellRange(range_string, min_row=start_row, min_col=start_column, max_row=end_row, max_col=end_column) if range_string else CellRange(min_row=start_row, min_col=start_column, max_row=end_row, max_col=end_column)
        if max(self.max_row, target.max_row) * max(self.max_column, target.max_col) > 500000: raise ValueError("Merged range exceeds CSV capacity")
        if any(not target.isdisjoint(existing) for existing in self.merged_cells.ranges):
            if target in self.merged_cells.ranges: return
            raise ValueError("Merged ranges must not overlap")
        for row in range(target.min_row, target.max_row + 1):
            for column in range(target.min_col, target.max_col + 1):
                if (row, column) != (target.min_row, target.min_col):
                    self._csv_shadow[row, column] = copy.copy(self.cell(row, column))
        super().merge_cells(str(target))

    def unmerge_cells(self, range_string=None, start_row=None, start_column=None, end_row=None, end_column=None):
        super().unmerge_cells(range_string, start_row, start_column, end_row, end_column)
        for key, cell in list(self._csv_shadow.items()):
            if key not in self._cells:
                self._cells[key] = cell
                del self._csv_shadow[key]

    def _structural_edit(self, *args, **kwargs):
        raise NotImplementedError("Use the CSV editor for row/column insertion, deletion or moving ranges; cell values and append() are supported here")

    insert_rows = delete_rows = insert_cols = delete_cols = move_range = _structural_edit


def _key(name): return "numberFormat" if name == "number_format" else name


class CsvWorkbook(Workbook):
    def __init__(self, client, payload, scope):
        super().__init__()
        self._client, self._scope = client, copy.deepcopy(scope)
        self._snapshot = CsvViewSnapshot.from_payload(payload)
        self._source = CsvSource(payload["content"], self._snapshot.delimiter)
        self._fonts[0] = Font(name="ui", sz=9.75, color=Color(auto=True))
        self._alignments[0] = Alignment(vertical="top", wrap_text=self._snapshot.wrap_text)
        self.remove(self.active)
        title = re.sub(r"[\\*?:/\[\]]", "_", PurePosixPath(self._snapshot.file_path).stem)[:31] or "CSV"
        ws = CsvWorksheet(self, title)
        self._add_sheet(ws)
        self._sheet = ws
        for row, values in enumerate(self._source.rows, 1):
            for column, value in enumerate(values, 1):
                cell = ws.cell(row, column, value)
                cell.data_type = "s"  # CSV formulas are literal text, as are identifiers and dates.
        for column in self._snapshot.columns.values():
            dim = ws.column_dimensions[get_column_letter(column.source_index + 1)]
            dim.width = (column.width - 5) / 7
            dim.hidden = column.hidden
        ws.sheet_format.defaultRowHeight = self._snapshot.row_height * 72 / 96
        for row, value in self._snapshot.row_dimensions.items():
            dim = ws.row_dimensions[int(row) + 1]
            dim.height, dim.hidden = value["height"], value["hidden"]
        if self._snapshot.frozen_columns:
            ws.freeze_panes = f"{get_column_letter(self._snapshot.frozen_columns + 1)}1"
        self._load_styles()
        self._load_conditionals()
        for merge in self._snapshot.merges:
            ws.merge_cells(start_row=merge["rows"][0] + 1, end_row=merge["rows"][1] + 1,
                start_column=merge["columns"][0] + 1, end_column=merge["columns"][1] + 1)
        self._remember()

    @property
    def revision(self): return self._snapshot.revision

    def _load_styles(self):
        ws = self._sheet
        self._original_excel = {}
        for rule in self._snapshot.styles:
            if rule.get("when"): continue  # Preserve dynamic rules; do not bake their current results into direct styles.
            if rule["id"].startswith("openpyxl-rows-"):
                apply_style(ws.row_dimensions[rule["rows"][0] + 1], rule["style"])
            elif rule["id"].startswith("openpyxl-columns-") and rule.get("columns"):
                column = self._snapshot.columns[rule["columns"][0]].source_index + 1
                apply_style(ws.column_dimensions[get_column_letter(column)], rule["style"])
            start, end = rule.get("rows", [0, max(0, ws.max_row - 1)])
            columns = [self._snapshot.columns[key].source_index + 1 for key in rule.get("columns", self._snapshot.column_order)]
            if not columns: continue
            if (end + 1) * max(columns) > 500000: raise ValueError("Style range exceeds CSV worksheet capacity")
            for row in range(start + 1, end + 2):
                for column in columns:
                    apply_style(ws.cell(row, column), rule["style"])
                    original = self._original_excel.setdefault((row, column), {})
                    if "color" in rule["style"]: original.setdefault("font", {})["color"] = rule["style"]["color"]
                    if "background" in rule["style"]: original.setdefault("fill", {})["fgColor"] = rule["style"]["background"]
                    original.update(copy.deepcopy(rule["style"].get("excel", {})))

    def _load_conditionals(self):
        from openpyxl.formatting.rule import CellIsRule
        operators = {"eq": "equal", "ne": "notEqual", "num_eq": "equal", "num_ne": "notEqual", "lt": "lessThan", "lte": "lessThanOrEqual", "gt": "greaterThan", "gte": "greaterThanOrEqual"}
        for rule in self._snapshot.styles:
            when = rule.get("when")
            if not rule["id"].startswith("openpyxl-cf-") or not when: continue
            if when["op"] not in operators or rule.get("columns") != [when["columnId"]] or "rows" not in rule:
                raise ValueError("Invalid persisted openpyxl conditional rule")
            column = get_column_letter(self._snapshot.columns[when["columnId"]].source_index + 1)
            first, last = rule["rows"]
            value = when["value"]
            formula = '"' + value.replace('"', '""') + '"' if isinstance(value, str) else str(value)
            cell = Cell(self._sheet, 1, 1)
            apply_style(cell, rule["style"])
            kwargs = {name: copy.copy(getattr(cell, name)) for name in ("font", "fill", "border") if name in rule["style"].get("excel", {})}
            legacy = rule["style"]
            if any(k in legacy for k in ("font", "size", "bold", "color")):
                from ._csv_excel import from_color
                kwargs["font"] = Font(**{name: value for name, value in {
                    "name": legacy.get("font"), "sz": legacy["size"] * 72 / 96 if "size" in legacy else None,
                    "b": legacy.get("bold"), "color": from_color(legacy["color"]) if "color" in legacy else None}.items() if value is not None})
            native = CellIsRule(operator=operators[when["op"]], formula=[formula], **kwargs)
            for name in ("alignment", "numberFormat"):
                if name in rule["style"].get("excel", {}):
                    if name == "alignment": native.dxf.alignment = decode_component(rule["style"]["excel"][name], name)
                    else:
                        from openpyxl.styles.numbers import NumberFormat
                        native.dxf.numFmt = NumberFormat(numFmtId=164, formatCode=rule["style"]["excel"][name])
            self._sheet.conditional_formatting.add(f"{column}{first + 1}:{column}{last + 1}", native)

    def _conditional_signature(self):
        from openpyxl.xml.functions import tostring
        return tuple((str(cf.sqref), tuple((tostring(rule.to_tree()), tostring(rule.dxf.to_tree()) if rule.dxf else None) for rule in rules)) for cf, rules in self._sheet.conditional_formatting._cf_rules.items())

    def _remember(self):
        ws = self._sheet
        style_cache = {}
        self._baseline = {}
        for key, cell in ws._cells.items():
            if isinstance(cell, MergedCell): continue
            signature = tuple(cell._style or ())
            if signature not in style_cache: style_cache[signature] = components(cell)
            self._baseline[key] = style_cache[signature]
        self._dimensions = {"rows": {k: copy.copy(v) for k, v in ws.row_dimensions.items()}, "columns": {k: copy.copy(v) for k, v in ws.column_dimensions.items()}}
        self._default_height = ws.sheet_format.defaultRowHeight
        self._cf_signature = self._conditional_signature()
        self._unsupported = self._unsupported_state()

    def _unsupported_state(self):
        ws = self._sheet
        from openpyxl.xml.functions import tostring
        return tuple(tostring(value.to_tree()) if hasattr(value, "to_tree") else repr(value) for value in (self.security, ws.protection, ws.auto_filter, ws.data_validations,
            ws.print_options, ws.page_setup, ws.page_margins, ws.print_area, ws.print_title_rows, ws.print_title_cols,
            self.defined_names, ws.sheet_state))

    def _check(self):
        ws = self._sheet
        if self.worksheets != [ws]: raise NotImplementedError("One CSV has exactly one worksheet")
        if ws.max_row * ws.max_column > 500000: raise ValueError("CSV exceeds 500,000 rectangular cells")
        if ws._charts or ws._images or ws.tables or self._unsupported != self._unsupported_state():
            raise NotImplementedError("CSV does not support workbook protection, printing, charts, images, tables, validation or Excel autofilters")
        for cell in ws._cells.values():
            if not isinstance(cell, MergedCell) and (cell.comment or cell.hyperlink or cell.protection != self._protections[0]):
                raise NotImplementedError("CSV does not support comments, hyperlinks or cell protection")

    def _style_rules(self):
        ws = self._sheet
        groups = {}
        # A fresh Cell uses the workbook's defaults without enlarging the worksheet.
        defaults = components(Cell(ws, 1, 1))
        for (row, column), cell in sorted(ws._cells.items()):
            if isinstance(cell, MergedCell): continue
            baseline = self._baseline.get((row, column), defaults)
            values = components(cell)
            delta = {_key(name): encode_component(values[name], name, self._original_excel.get((row, column), {}).get(_key(name)))
                for name in COMPONENTS if values[name] != baseline[name]}
            if delta:
                key = json.dumps(delta, sort_keys=True, ensure_ascii=False, separators=(",", ":"))
                groups.setdefault(key, {}).setdefault(row, []).append(column)
        rules = []
        for encoded, rows in groups.items():
            # Collapse equal adjacent rows and contiguous columns into rectangles.
            ranges = []
            for row, columns in rows.items():
                if ranges and ranges[-1][1] + 1 == row and ranges[-1][2] == columns: ranges[-1][1] = row
                else: ranges.append([row, row, columns])
            for first, last, columns in ranges:
                target = {"rows": [first - 1, last - 1], "columns": [{"source_index": c - 1} for c in columns]}
                signature = json.dumps([target, sorted(json.loads(encoded))], sort_keys=True)
                rules.append({"id": "openpyxl-" + hashlib.sha256(signature.encode()).hexdigest()[:24], **target, "style": {"excel": json.loads(encoded)}})
        for axis, dimensions in (("rows", ws.row_dimensions), ("columns", ws.column_dimensions)):
            for index, dim in dimensions.items():
                baseline_dim = self._dimensions[axis].get(index)
                baseline = components(baseline_dim) if baseline_dim else defaults
                current = components(dim)
                delta = {_key(name): encode_component(current[name], name) for name in COMPONENTS if current[name] != baseline[name]}
                if dim.outlineLevel or dim.collapsed: raise NotImplementedError("CSV does not support row/column outlines")
                if delta:
                    target = {"rows": [index - 1, index - 1]} if axis == "rows" else {"columns": [{"source_index": c - 1} for c in range(dim.min or column_index_from_string(index), (dim.max or column_index_from_string(index)) + 1)]}
                    identifier = f"openpyxl-{axis}-{index}"
                    old = next((r["style"].get("excel", {}) for r in self._snapshot.styles if r["id"] == identifier), {})
                    rules.append({"id": identifier, **target, "style": {"excel": {**old, **delta}}})
        return rules

    def _conditional_rules(self):
        ws = self._sheet
        if self._cf_signature == self._conditional_signature(): return [], []
        result = []
        seen = 0
        operators = {"equal": "eq", "notEqual": "ne", "lessThan": "lt", "lessThanOrEqual": "lte", "greaterThan": "gt", "greaterThanOrEqual": "gte"}
        for cf, entries in ws.conditional_formatting._cf_rules.items():
            for rule in entries:
                seen += 1
                if rule.type != "cellIs" or rule.operator not in operators or len(rule.formula) != 1 or rule.stopIfTrue:
                    raise NotImplementedError("CSV conditional formatting supports CellIsRule with one literal comparison; formulas, stopIfTrue and color scales are not supported")
                literal = str(rule.formula[0])
                if re.fullmatch(r'"(?:[^"]|"")*"', literal): value = literal[1:-1].replace('""', '"')
                elif re.fullmatch(r"[+-]?(?:\d+(?:\.\d*)?|\.\d+)(?:e[+-]?\d+)?", literal, re.I): value = float(literal)
                else: raise NotImplementedError("CSV conditional comparisons require a literal value")
                dxf = rule.dxf
                if dxf is None: raise ValueError("Conditional rule has no formatting")
                style = {_key(name): encode_component(copy.copy(getattr(dxf, name)), name) for name in ("fill", "border", "alignment") if getattr(dxf, name) is not None}
                legacy = {}
                if dxf.font is not None:
                    f = dxf.font
                    specified = {child.tag.rsplit("}", 1)[-1] for child in f.to_tree()}
                    if f.i or f.u or f.strike or f.vertAlign or f.outline or f.shadow:
                        raise NotImplementedError("Conditional fonts support name, size, bold and color; use direct styles for other font effects")
                    if f.name is not None: legacy["font"] = f.name
                    if f.sz is not None: legacy["size"] = round(f.sz * 96 / 72)
                    if "b" in specified: legacy["bold"] = f.b
                    if f.color is not None: legacy["color"] = to_color(f.color, "default")
                if dxf.numFmt: style["numberFormat"] = dxf.numFmt.formatCode
                if not style and not legacy: raise ValueError("Conditional rule has no supported formatting")
                operator = {"equal": "num_eq", "notEqual": "num_ne"}.get(rule.operator, operators[rule.operator]) if not isinstance(value, str) else operators[rule.operator]
                for area in sorted(cf.sqref.ranges, key=str):
                    for column in range(area.min_col, area.max_col + 1):
                        target = {"source_index": column - 1}
                        result.append({"id": "openpyxl-cf-" + hashlib.sha256(f"{seen}:{area}:{column}".encode()).hexdigest()[:24],
                            "rows": [area.min_row - 1, area.max_row - 1], "columns": [target],
                            "when": {"target": target, "op": operator, "value": value}, "style": {**legacy, **({"excel": style} if style else {})}})
        ids = {r["id"] for r in result}
        removed = [r["id"] for r in self._snapshot.styles if r["id"].startswith("openpyxl-cf-") and r["id"] not in ids]
        return result, removed

    async def save(self, filename=None):
        """Save to the original CSV through Locus. Revisions / batching are internal."""
        if filename is not None and str(filename).replace("\\", "/") != self._snapshot.file_path:
            raise ValueError("Save this workbook to its original CSV; omit filename")
        self._check()
        ws = self._sheet
        conditional, removed = self._conditional_rules()
        rules = self._style_rules() + conditional
        patch = {}
        count = max(ws.max_column, max((column_index_from_string(c) for c in ws.column_dimensions), default=1))
        if count > max((c.source_index + 1 for c in self._snapshot.columns.values()), default=0): patch["column_count"] = count
        if rules or removed:
            modified = {r["id"] for r in rules}
            patch["styles"] = {"upsert": rules, "remove": removed, "order": [r["id"] for r in self._snapshot.styles if r["id"] not in modified and r["id"] not in removed] + [r["id"] for r in rules]}
        columns = []
        for label, dim in ws.column_dimensions.items():
            old = self._dimensions["columns"].get(label)
            if old is None or old.width != dim.width or old.hidden != dim.hidden:
                width = round(dim.width * 7 + 5)
                if not 48 <= width <= 2000: raise ValueError("CSV column width must convert to 48..2000 pixels")
                for column in range(dim.min or column_index_from_string(label), (dim.max or column_index_from_string(label)) + 1):
                    columns.append({"target": {"source_index": column - 1}, "width": width, "hidden": bool(dim.hidden)})
        if columns: patch["columns"] = columns
        dimensions = {str(row - 1): {"height": float(dim.height if dim.height is not None else ws.sheet_format.defaultRowHeight), "hidden": bool(dim.hidden)}
            for row, dim in ws.row_dimensions.items() if dim.height is not None or dim.hidden}
        if dimensions != self._snapshot.row_dimensions: patch["row_dimensions"] = dimensions
        if ws.sheet_format.defaultRowHeight != self._default_height:
            height = round(ws.sheet_format.defaultRowHeight * 96 / 72)
            if not 20 <= height <= 120: raise ValueError("CSV default row height must convert to 20..120 pixels; use row_dimensions for other heights")
            patch["row_height"] = height
        row, col = coordinate_to_tuple(ws.freeze_panes or "A1")
        if row != 1: raise NotImplementedError("CSV supports freezing leading columns; freezing rows is not supported")
        if col - 1 != self._snapshot.frozen_columns: patch["frozen_columns"] = col - 1
        merges = [{"rows": [r.min_row - 1, r.max_row - 1], "columns": [r.min_col - 1, r.max_col - 1]} for r in sorted(ws.merged_cells.ranges, key=lambda r: (r.min_row, r.min_col))]
        if merges != self._snapshot.merges: patch["merges"] = merges
        values = {key: text_value(cell.value) for key, cell in {**ws._cells, **ws._csv_shadow}.items() if not isinstance(cell, MergedCell)}
        content = self._source.edit(values, minimum_rows=ws._csv_appended_extent)
        if not patch and content == self._source.content: return self._snapshot
        payload = {"workspaceRef": self._scope, "filePath": self._snapshot.file_path, "expectedRevision": self.revision,
            "patch": patch, "sessionId": self._client.current_session_id,
            "executionDelegation": os.environ.get("LOCUS_SDK_EXECUTION_DELEGATION")}
        if content != self._source.content: payload["content"] = content
        result = await self._client.rpc("csv.save_workbook", payload)
        self._snapshot = CsvViewSnapshot.from_payload(result)
        self._source = CsvSource(content, self._snapshot.delimiter)
        self._remember()
        return self._snapshot

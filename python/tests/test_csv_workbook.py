import copy
import os
import unittest
from unittest.mock import AsyncMock, patch

import locus
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, GradientFill, NamedStyle, PatternFill, Side
from openpyxl.formatting.rule import CellIsRule

from locus._csv_codec import CsvSource


def payload(content="id,name,value\r\n001,one,0.125\r\n002,two,1234.50\r\n"):
    return {"filePath": "items.csv", "revision": "v1", "rowCount": 3, "columnCount": 3, "delimiter": ",", "content": content,
        "view": {"schema": "locus.csv-view.v1", "headerRows": 1, "rowHeight": 28, "wrapText": False, "frozenColumns": 0,
            "columns": {f"c{i}": {"sourceIndex": i, "header": title, "width": 140} for i, title in enumerate(["id", "name", "value"])},
            "columnOrder": ["c0", "c1", "c2"]}}


class CsvWorkbookTests(unittest.IsolatedAsyncioTestCase):
    def setUp(self):
        self.client = locus.Client(base_url="http://127.0.0.1/sdk", token="test", current_session_id="session")
        self.client.rpc = AsyncMock(return_value=payload())
        self.scope = locus.WorkspaceRef("test")

    async def load(self, data=None):
        self.client.rpc.return_value = data or payload()
        book = await self.client.csv.load_workbook("items.csv", workspace_ref=self.scope)
        self.client.rpc.reset_mock()
        return book

    async def test_real_openpyxl_objects_named_styles_ranges_and_noop_save(self):
        book = await self.load()
        self.assertIsInstance(book, Workbook)
        ws = book.active
        self.assertEqual(ws["A2"].value, "001")
        self.assertEqual(tuple(ws.iter_rows(min_row=2, max_row=2, values_only=True)), (("001", "one", "0.125"),))
        await book.save()
        self.client.rpc.assert_not_awaited()
        header = NamedStyle(name="header", font=Font(name="Arial", size=12.5, bold=True, italic=True, color="FF123456"), fill=PatternFill("solid", fgColor="ABCDEF"))
        book.add_named_style(header)
        for row in ws["A1:C1"]:
            for cell in row: cell.style = "header"
        ws["C2"].number_format = "0.0%"
        ws["B2"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True, indent=1, text_rotation=30)
        ws["B2"].border = Border(left=Side(style="thin", color="123456"), bottom=Side(style="double", color="FF0000"))
        ws.column_dimensions["B"].width = 24
        ws.row_dimensions[1].height = 32
        ws.row_dimensions[3].hidden = True
        ws.freeze_panes = "B1"
        await book.save()
        method, request = self.client.rpc.call_args.args
        self.assertEqual(method, "csv.save_workbook")
        self.assertNotIn("content", request)
        self.assertEqual(request["expectedRevision"], "v1")
        update = request["patch"]
        self.assertEqual(update["columns"][0]["width"], 173)
        self.assertEqual(update["row_dimensions"]["0"], {"height": 32, "hidden": False})
        rules = update["styles"]["upsert"]
        self.assertEqual(len(rules), 3)
        font = rules[0]["style"]["excel"]["font"]
        self.assertEqual((font["size"], font["color"], font["italic"]), (12.5, "#123456", True))
        self.assertEqual(len(rules[0]["columns"]), 3)

    async def test_csv_bytes_and_literals_are_preserved_when_only_one_field_changes(self):
        content = '\ufeffid,name,value\r\n001,"one\r\ntwo",0.125\r\n002,"two",1234.50'
        book = await self.load(payload(content))
        book.active["C2"] = 0.5
        await book.save()
        request = self.client.rpc.call_args.args[1]
        self.assertEqual(request["content"], content.replace("0.125", "0.5"))
        self.assertEqual(request["patch"], {})
        self.client.rpc.reset_mock()
        await book.save()
        self.client.rpc.assert_not_awaited()

    async def test_csv_strings_are_not_truncated_by_excel_cell_limits(self):
        long_value = "a" * 40000 + "\x01"
        book = await self.load(payload("id,name,value\n001," + long_value + ",=literal\n"))
        self.assertEqual(book.active["B2"].value, long_value)
        self.assertEqual(book.active["C2"].data_type, "s")
        book.active["B2"].font = Font(bold=True)
        await book.save()
        self.assertNotIn("content", self.client.rpc.call_args.args[1])
        book.active.append(["002", long_value, "=more"])
        await book.save()
        self.assertIn(long_value, self.client.rpc.call_args.args[1]["content"])

    async def test_explicit_blank_rows_can_be_appended_without_styling_creating_data(self):
        book = await self.load()
        book.active.append(["", "", ""])
        await book.save()
        self.assertEqual(self.client.rpc.call_args.args[1]["content"], payload()["content"] + ",,")

    async def test_merge_and_unmerge_preserve_covered_csv_values(self):
        book = await self.load()
        ws = book.active
        ws.merge_cells("A2:B2")
        await book.save()
        self.assertNotIn("content", self.client.rpc.call_args.args[1])
        ws.unmerge_cells("A2:B2")
        self.assertEqual(ws["B2"].value, "one")
        with self.assertRaises(ValueError):
            ws.merge_cells("A1:C2")
            ws.merge_cells("B2:C3")

    async def test_saves_compress_ranges_and_conflicts_keep_pending_edits(self):
        book = await self.load()
        for row in book.active.iter_rows(min_row=1, max_row=100, max_col=3):
            for cell in row: cell.font = Font(bold=True)
        self.client.rpc.side_effect = locus.LocusRpcError("csv.revision_changed")
        for _ in range(2):
            with self.assertRaises(locus.LocusRpcError): await book.save()
            request = self.client.rpc.call_args.args[1]
            self.assertEqual(request["expectedRevision"], "v1")
            self.assertNotIn("content", request)
            self.assertEqual(len(request["patch"]["styles"]["upsert"]), 1)
            self.assertEqual(request["patch"]["styles"]["upsert"][0]["rows"], [0, 99])

    async def test_existing_styles_are_readable_and_unchanged_rules_are_preserved(self):
        data = payload()
        data["view"]["schema"] = "locus.csv-view.v2"
        data["view"]["styles"] = [{"id": "header", "rows": [0, 0], "style": {"bold": True, "color": "accent"}},
            {"id": "warn", "when": {"columnId": "c2", "op": "gt", "value": 100}, "style": {"background": "warning-soft"}}]
        book = await self.load(data)
        self.assertTrue(book.active["A1"].font.bold)
        book.active["A1"].font = Font(bold=False)
        await book.save()
        request = self.client.rpc.call_args.args[1]
        self.assertEqual(request["patch"]["styles"]["order"][:2], ["header", "warn"])
        self.assertFalse(request["patch"]["styles"]["upsert"][0]["style"]["excel"]["font"]["bold"])

    async def test_unsupported_features_fail_before_any_write(self):
        for change in [lambda ws: setattr(ws["A1"], "fill", GradientFill(stop=("FFFFFF", "000000"))),
            lambda ws: setattr(ws.protection, "sheet", True), lambda ws: setattr(ws, "freeze_panes", "B2")]:
            book = await self.load()
            change(book.active)
            with self.assertRaises(NotImplementedError): await book.save()
            self.client.rpc.assert_not_awaited()

    async def test_native_conditional_formatting_maps_to_source_comparisons(self):
        book = await self.load()
        book.active.conditional_formatting.add("C2:C100", CellIsRule(operator="greaterThan", formula=[10], fill=PatternFill("solid", fgColor="FFCCAA")))
        await book.save()
        rule = self.client.rpc.call_args.args[1]["patch"]["styles"]["upsert"][0]
        self.assertEqual(rule["when"], {"target": {"source_index": 2}, "op": "gt", "value": 10})
        self.assertEqual(rule["rows"], [1, 99])

    async def test_conditional_rules_reload_update_and_remove_without_losing_other_rules(self):
        data = payload()
        data["view"]["schema"] = "locus.csv-view.v4"
        data["view"]["styles"] = [{"id": "openpyxl-cf-old", "rows": [1, 99], "columns": ["c2"],
            "when": {"columnId": "c2", "op": "num_eq", "value": 10}, "style": {"bold": True}}]
        book = await self.load(data)
        self.assertEqual(len(book.active.conditional_formatting), 1)
        await book.save()
        self.client.rpc.assert_not_awaited()
        rules = book.active.conditional_formatting["C2:C100"]
        rules[0].dxf.font = Font(bold=True, color="FF0000")
        await book.save()
        style = self.client.rpc.call_args.args[1]["patch"]["styles"]["upsert"][0]["style"]
        self.assertEqual(style, {"bold": True, "color": "#FF0000"})
        book = await self.load(data)
        del book.active.conditional_formatting["C2:C100"]
        await book.save()
        self.assertEqual(self.client.rpc.call_args.args[1]["patch"]["styles"]["remove"], ["openpyxl-cf-old"])

    async def test_row_style_updates_preserve_other_components_and_clearing_a_cell_persists(self):
        data = payload()
        data["view"]["schema"] = "locus.csv-view.v4"
        data["view"]["styles"] = [{"id": "openpyxl-rows-1", "rows": [0, 0], "style": {"excel": {"font": {
            "name": "Arial", "size": 12, "bold": True, "italic": False, "strike": False, "underline": "none", "color": "default", "vertAlign": "baseline"}}}}]
        book = await self.load(data)
        self.assertTrue(book.active.row_dimensions[1].font.bold)
        book.active.row_dimensions[1].fill = PatternFill("solid", fgColor="ABCDEF")
        del book.active["B2"]
        await book.save()
        request = self.client.rpc.call_args.args[1]
        self.assertIn("001,,0.125", request["content"])
        row = next(rule for rule in request["patch"]["styles"]["upsert"] if rule["id"] == "openpyxl-rows-1")
        self.assertTrue(row["style"]["excel"]["font"]["bold"])
        self.assertEqual(row["style"]["excel"]["fill"]["fgColor"], "#ABCDEF")


class CsvCodecTests(unittest.TestCase):
    def test_special_records_and_growth(self):
        for text, delimiter in [('a,b,\r\n"x,y","z""q",\n\n', ','), ('\ufeffa;b\r001;"a\rb"', ';'), ('a\tb\n1\t2\n', '\t'), ('', ',')]:
            source = CsvSource(text, delimiter)
            self.assertEqual(source.edit({}), text)
            for row, values in enumerate(source.rows, 1):
                for column, value in enumerate(values, 1): self.assertEqual(source.edit({(row, column): value}), text)
        source = CsvSource('a,b\r\n001,"two"', ',')
        self.assertEqual(source.edit({(2, 2): 'new"value'}), 'a,b\r\n001,"new""value"')
        self.assertEqual(source.edit({(4, 3): 'tail'}), 'a,b\r\n001,"two"\r\n\r\n,,tail')


if __name__ == "__main__": unittest.main()
